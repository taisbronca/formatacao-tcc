#!/usr/bin/env python3
"""Structural preflight for USP/Esalq TCC DOCX files.

The output contains deterministic observations, not final compliance decisions.
Page count and visual layout still require rendering and human inspection.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
import unicodedata
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


PT_TOLERANCE = 0.2
REQUIRED_TCC = [
    "resumo|sumário executivo",
    "palavras-chave",
    "introdução|considerações iniciais",
    "metodologia|material e métodos|implementação de algoritmo de machine learning|implementação de algoritmos de machine learning",
    "resultados e discussão",
    "conclusão|conclusões|considerações finais",
    "referências",
]
KNOWN_HEADINGS = {
    "resumo",
    "sumário executivo",
    "abstract",
    "resumen",
    "palavras-chave",
    "keywords",
    "palabras clave",
    "introdução",
    "considerações iniciais",
    "metodologia",
    "material e métodos",
    "implementação de algoritmos de machine learning",
    "implementação de algoritmo de machine learning",
    "resultados preliminares",
    "resultados e discussão",
    "conclusão",
    "conclusões",
    "considerações finais",
    "agradecimentos",
    "referências",
}


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return re.sub(r"\s+", " ", text.strip().lower())


KNOWN_HEADINGS_NORM = {norm(x) for x in KNOWN_HEADINGS}


def cm(value) -> float | None:
    return None if value is None else round(value.cm, 3)


def pt(value) -> float | None:
    return None if value is None else round(value.pt, 2)


def color_of(run) -> str | None:
    rgb = run.font.color.rgb
    return None if rgb is None else str(rgb)


def effective_font_name(run, paragraph) -> str | None:
    if run.font.name:
        return run.font.name
    if paragraph.style and paragraph.style.font.name:
        return paragraph.style.font.name
    return None


def effective_font_size(run, paragraph) -> float | None:
    if run.font.size:
        return pt(run.font.size)
    if paragraph.style and paragraph.style.font.size:
        return pt(paragraph.style.font.size)
    return None


def alignment_name(value) -> str:
    names = {
        None: "herdado",
        WD_ALIGN_PARAGRAPH.LEFT: "esquerda",
        WD_ALIGN_PARAGRAPH.CENTER: "centralizado",
        WD_ALIGN_PARAGRAPH.RIGHT: "direita",
        WD_ALIGN_PARAGRAPH.JUSTIFY: "justificado",
        WD_ALIGN_PARAGRAPH.DISTRIBUTE: "distribuído",
    }
    return names.get(value, str(value))


def line_spacing_value(paragraph):
    value = paragraph.paragraph_format.line_spacing
    if value is None:
        return None
    if hasattr(value, "pt"):
        return {"points": pt(value)}
    if isinstance(value, (int, float)):
        return round(float(value), 3)
    return str(value)


def paragraph_runs(paragraph):
    return [run for run in paragraph.runs if run.text.strip()]


def paragraph_profile(paragraph) -> dict:
    runs = paragraph_runs(paragraph)
    fonts = Counter(filter(None, (effective_font_name(run, paragraph) for run in runs)))
    sizes = Counter(filter(None, (effective_font_size(run, paragraph) for run in runs)))
    colors = Counter(filter(None, (color_of(run) for run in runs)))
    return {
        "style": paragraph.style.name if paragraph.style else None,
        "alignment": alignment_name(paragraph.alignment),
        "line_spacing": line_spacing_value(paragraph),
        "first_line_indent_cm": cm(paragraph.paragraph_format.first_line_indent),
        "space_before_pt": pt(paragraph.paragraph_format.space_before),
        "space_after_pt": pt(paragraph.paragraph_format.space_after),
        "fonts": dict(fonts),
        "sizes_pt": {str(k): v for k, v in sizes.items()},
        "colors": dict(colors),
        "all_bold": bool(runs) and all(run.bold is True for run in runs),
        "any_italic": any(run.italic is True for run in runs),
    }


def find_heading_indices(paragraphs) -> dict[str, int]:
    found = {}
    for idx, paragraph in enumerate(paragraphs):
        value = norm(paragraph.text).rstrip(":")
        if value in KNOWN_HEADINGS_NORM:
            found.setdefault(value, idx)
    return found


def section_text(paragraphs, start_names: set[str], stop_names: set[str]) -> tuple[int | None, list]:
    start = None
    collected = []
    for idx, paragraph in enumerate(paragraphs):
        value = norm(paragraph.text).rstrip(":")
        if start is None and value in start_names:
            start = idx
            continue
        if start is not None:
            if value in stop_names or any(value.startswith(stop + ":") for stop in stop_names):
                break
            if paragraph.text.strip():
                collected.append(paragraph)
    return start, collected


def has_page_field(section) -> bool:
    for paragraph in section.footer.paragraphs:
        xml = paragraph._p.xml.upper()
        if "PAGE" in xml and ("INSTRTEXT" in xml or "FLDSIMPLE" in xml):
            return True
    return False


def table_observation(table, index: int) -> dict:
    bold_cells = 0
    shaded_cells = 0
    numeric_alignment = Counter()
    total_cells = 0
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            total_cells += 1
            tc_pr = cell._tc.get_or_add_tcPr()
            if tc_pr.find(qn("w:shd")) is not None:
                shaded_cells += 1
            text = " ".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
            runs = [run for p in cell.paragraphs for run in paragraph_runs(p)]
            if r_idx > 0 and runs and any(run.bold is True for run in runs):
                bold_cells += 1
            if r_idx > 0 and re.fullmatch(r"[\sR$%+\-–—\d.,()/]+", text or "x"):
                for paragraph in cell.paragraphs:
                    if paragraph.text.strip():
                        numeric_alignment[alignment_name(paragraph.alignment)] += 1
    header_repeat = False
    if table.rows:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        header_repeat = tr_pr.find(qn("w:tblHeader")) is not None
    return {
        "table": index,
        "rows": len(table.rows),
        "columns": len(table.columns),
        "cells": total_cells,
        "body_cells_with_bold": bold_cells,
        "cells_with_shading": shaded_cells,
        "numeric_alignment": dict(numeric_alignment),
        "first_row_repeats": header_repeat,
    }


def audit(path: Path, stage: str) -> dict:
    doc = Document(path)
    paragraphs = list(doc.paragraphs)
    nonempty = [(idx, p) for idx, p in enumerate(paragraphs) if p.text.strip()]
    headings = find_heading_indices(paragraphs)

    sections = []
    for idx, section in enumerate(doc.sections, 1):
        orientation = "paisagem" if section.orientation == WD_ORIENT.LANDSCAPE else "retrato"
        sections.append(
            {
                "section": idx,
                "orientation": orientation,
                "page_width_cm": cm(section.page_width),
                "page_height_cm": cm(section.page_height),
                "margins_cm": {
                    "top": cm(section.top_margin),
                    "bottom": cm(section.bottom_margin),
                    "left": cm(section.left_margin),
                    "right": cm(section.right_margin),
                },
                "header_text": " | ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip()),
                "footer_text": " | ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip()),
                "footer_has_page_field": has_page_field(section),
            }
        )

    paragraph_issues = []
    for idx, paragraph in nonempty:
        profile = paragraph_profile(paragraph)
        fonts = {name.lower() for name in profile["fonts"]}
        bad_fonts = sorted(name for name in fonts if name != "arial")
        bad_sizes = sorted(float(size) for size in profile["sizes_pt"] if abs(float(size) - 11.0) > PT_TOLERANCE)
        nonblack = [color for color in profile["colors"] if color != "000000"]
        if bad_fonts or bad_sizes or nonblack:
            paragraph_issues.append(
                {
                    "paragraph": idx + 1,
                    "text_excerpt": paragraph.text[:120],
                    "non_arial_fonts": bad_fonts,
                    "non_11pt_sizes": bad_sizes,
                    "explicit_nonblack_colors": nonblack,
                    "note": "Exceções legítimas (cabeçalho, rodapé, endereços, notas) exigem confirmação contextual.",
                }
            )

    abstract_names = {"resumo", "sumario executivo"}
    keyword_names = {"palavras-chave", "palavras chave", "keywords", "palabras clave"}
    abstract_start, abstract_paragraphs = section_text(paragraphs, abstract_names, keyword_names)
    abstract_text = " ".join(p.text.strip() for p in abstract_paragraphs)
    abstract_words = re.findall(r"\b[\wÀ-ÿ'-]+\b", abstract_text)
    citation_like = re.findall(r"\([A-ZÀ-Ü][^()]{0,80},\s*(?:19|20)\d{2}[a-z]?\)", abstract_text)

    keyword_text = ""
    if abstract_start is not None:
        for paragraph in paragraphs[abstract_start + 1 :]:
            text = paragraph.text.strip()
            normalized = norm(text)
            if normalized.startswith("palavras-chave") or normalized.startswith("palavras chave"):
                keyword_text = re.sub(r"^[^:]+:\s*", "", text)
                break
    keywords = [item.strip(" .") for item in keyword_text.split(";") if item.strip(" .")]

    all_text_norm = "\n".join(norm(p.text) for p in paragraphs if p.text.strip())
    required = []
    if stage == "tcc":
        for spec in REQUIRED_TCC:
            variants = [norm(item) for item in spec.split("|")]
            required.append({"requirement": spec, "present": any(value in all_text_norm for value in variants)})

    candidate_title = None
    candidates = []
    for idx, paragraph in nonempty[:20]:
        if norm(paragraph.text) not in KNOWN_HEADINGS_NORM and len(paragraph.text.split()) >= 3:
            profile = paragraph_profile(paragraph)
            score = int(profile["alignment"] == "centralizado") + int(profile["all_bold"])
            candidates.append((score, idx, paragraph))
    if candidates:
        _, idx, paragraph = max(candidates, key=lambda item: (item[0], -item[1]))
        candidate_title = {
            "paragraph": idx + 1,
            "text": paragraph.text.strip(),
            "word_count": len(re.findall(r"\b[\wÀ-ÿ'-]+\b", paragraph.text)),
            "profile": paragraph_profile(paragraph),
            "note": "Candidato heurístico; confirme visualmente na folha de rosto.",
        }

    pictures = len(doc.inline_shapes)
    figure_caption_count = sum(bool(re.match(r"^\s*Figura\s+\d+\.", p.text, re.I)) for p in paragraphs)
    table_title_count = sum(bool(re.match(r"^\s*Tabela\s+\d+\.", p.text, re.I)) for p in paragraphs)

    return {
        "file": str(path.resolve()),
        "stage": stage,
        "limitations": [
            "A contagem de páginas e a conformidade visual exigem renderização integral.",
            "Fonte/tamanho herdados podem não aparecer como propriedades explícitas.",
            "Detecção de título, resumo, palavras-chave e seções é heurística e requer confirmação.",
            "Redação, tempo verbal, integridade científica e ética exigem revisão humana.",
        ],
        "document": {
            "paragraphs": len(paragraphs),
            "tables": len(doc.tables),
            "inline_shapes": pictures,
            "sections": sections,
        },
        "required_sections": required,
        "detected_headings": [
            {"heading": name, "paragraph": idx + 1}
            for name, idx in sorted(headings.items(), key=lambda item: item[1])
        ],
        "title_candidate": candidate_title,
        "abstract": {
            "found": abstract_start is not None,
            "paragraph_count": len(abstract_paragraphs),
            "word_count": len(abstract_words),
            "citation_like_patterns": citation_like,
        },
        "keywords": {"raw": keyword_text, "count": len(keywords), "items": keywords},
        "formatting_observations": paragraph_issues,
        "tables": [table_observation(table, idx) for idx, table in enumerate(doc.tables, 1)],
        "figures": {
            "inline_shapes": pictures,
            "figure_caption_paragraphs": figure_caption_count,
            "table_title_paragraphs": table_title_count,
            "note": "Compare objetos, menções, legendas e fontes manualmente; gráficos flutuantes podem não contar como inline_shapes.",
        },
        "page_count": {"value": None, "status": "render_required"},
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Auditoria estrutural preliminar de DOCX conforme o manual USP/Esalq.")
    parser.add_argument("docx", type=Path)
    parser.add_argument("--stage", choices=["tcc", "resultados-preliminares", "projeto"], default="tcc")
    parser.add_argument("--json", type=Path, dest="json_path", help="Grava o relatório JSON neste caminho.")
    args = parser.parse_args()

    if not args.docx.exists():
        parser.error(f"arquivo não encontrado: {args.docx}")
    if args.docx.suffix.lower() != ".docx":
        parser.error("o arquivo deve ter extensão .docx")

    try:
        result = audit(args.docx, args.stage)
    except Exception as exc:
        print(json.dumps({"error": str(exc), "file": str(args.docx)}, ensure_ascii=False, indent=2), file=sys.stderr)
        return 2

    output = json.dumps(result, ensure_ascii=False, indent=2)
    if args.json_path:
        args.json_path.parent.mkdir(parents=True, exist_ok=True)
        args.json_path.write_text(output + "\n", encoding="utf-8")
    print(output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
